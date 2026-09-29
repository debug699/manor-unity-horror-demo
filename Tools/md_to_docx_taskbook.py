from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import re

src=Path(r'D:\游戏制作软件\参考图\下一步单独生图提示词_已检查布局版.md')
out=Path(r'C:\Users\李\Desktop\《庄园》下一步单独生图提示词_已检查布局版.docx')
lines=src.read_text(encoding='utf-8').splitlines()
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles
for name,size,color in [('Normal',10.5,'243B53'),('Heading 1',16,'7B1E24'),('Heading 2',13,'334E68'),('Heading 3',11,'52606D')]:
    s=styles[name]; s.font.name='Aptos'; s._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
def add_run(p,text,bold=False,code=False):
    r=p.add_run(text); r.bold=bold; r.font.name='Consolas' if code else 'Aptos'; r._element.rPr.rFonts.set(qn('w:ascii'),r.font.name); r._element.rPr.rFonts.set(qn('w:hAnsi'),r.font.name); r.font.size=Pt(9 if code else 10.5); return r
def inline(p,text):
    parts=re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text)
    for x in parts:
        if x.startswith('`') and x.endswith('`'): add_run(p,x[1:-1],code=True)
        elif x.startswith('**') and x.endswith('**'): add_run(p,x[2:-2],bold=True)
        else: add_run(p,x)
def table(rows):
    cols=max(len(r) for r in rows); t=doc.add_table(rows=0,cols=cols); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.LEFT
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            p=cells[i].paragraphs[0]; inline(p,val)
            if ri==0:
                for run in p.runs: run.bold=True; run.font.color.rgb=RGBColor.from_string('243B53')
    doc.add_paragraph('')
i=0; in_code=False; code=[]; tbl=[]
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        if not in_code:
            in_code=True; code=[]
        else:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.15); p.paragraph_format.right_indent=Inches(.1); add_run(p,'\n'.join(code),code=True); in_code=False
        i+=1; continue
    if in_code: code.append(line); i+=1; continue
    if line.startswith('|'):
        if set(line.replace('|','').replace('-','').replace(':','').strip())==set(): i+=1; continue
        tbl.append([x.strip() for x in line.strip('|').split('|')]); i+=1
        if i>=len(lines) or not lines[i].startswith('|'):
            table(tbl); tbl=[]
        continue
    m=re.match(r'^(#{1,3})\s+(.*)',line)
    if m:
        p=doc.add_paragraph(style=f'Heading {len(m.group(1))}'); add_run(p,m.group(2),bold=True); i+=1; continue
    m=re.match(r'^\s*(\d+)\.\s+(.*)',line)
    if m:
        p=doc.add_paragraph(style='List Number'); inline(p,m.group(2)); i+=1; continue
    if line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); inline(p,line[2:]); i+=1; continue
    if line.strip():
        p=doc.add_paragraph(); inline(p,line)
    i+=1
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; add_run(footer,'《庄园》AI 美术与 Unity 制作任务书 | 版本 1.0')
doc.save(out); print(out)
